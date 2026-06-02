from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def find_project_root(start_dir: Path) -> Path:
    for path in [start_dir, *start_dir.parents]:
        if (path / "requirements.txt").exists() and (path / "source").exists():
            return path
        if path.name == "organized_workspace_20260601":
            return path
    return start_dir


ROOT = find_project_root(Path(__file__).resolve().parent)
SRC = ROOT / "第二十一届研电赛-技术论文模板-0508v1.docx"
OUT = ROOT / "第二十一届研电赛-技术论文模板-0508v1_USRP编译码填充.docx"


def set_run_font(run, size_pt: float | None = None, bold: bool | None = None):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if size_pt:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold


def set_para_text(paragraph, text: str, size_pt: float | None = None, bold: bool | None = None):
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run, size_pt=size_pt, bold=bold)


def paragraph_after(paragraph, text: str = "", style: str | None = None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = paragraph._parent.add_paragraph()
    p._p = new_p
    if style:
        p.style = style
    if text:
        set_para_text(p, text, size_pt=10.5)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Pt(21)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(4)
    return p


def move_table_after(doc: Document, anchor_paragraph, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r == 0 or c != len(row) - 1 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(text))
            set_run_font(run, size_pt=9, bold=(r == 0))
            if r == 0:
                tc_pr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement("w:shd")
                shd.set(qn("w:fill"), "D9EAF7")
                tc_pr.append(shd)
    anchor_paragraph._p.addnext(table._tbl)
    return table


def find_paragraph(doc: Document, exact: str):
    for p in doc.paragraphs:
        if p.text.strip() == exact:
            return p
    raise ValueError(f"paragraph not found: {exact}")


def remove_instruction_paragraphs(doc: Document):
    markers = [
        "一句话说应用场景，一句话说核心方法，一句话说实测效果。别写空泛口号，要出现关键指标，比如精度、功耗、延迟、体积、成本、吞吐率等。",
        "写清楚你解决了什么痛点、现有方法有什么不足、你的作品为什么值得做。这里不要写成综述，重点是把“需求”和“评价指标”立住。",
        "建议列 3 个创新点，每个都按“问题-方法-效果”写。\n比如：针对 XX 场景中 XX 延迟高的问题，提出 XX 架构，使端到端延迟降低到 XX ms。",
        "给出一段各自的难点与创新，后续整合与优化：",
        "BSY：通用译码器",
        "GC：信道预测与自适应",
        "XL：软硬件协同系统仿真",
        "这里放系统框图、模块划分、数据流/控制流。说明为什么选这个方案，而不是其他方案。",
        "这里讲如何讲通用译码算法在GPU架构上实现并部署",
        "这里讲如何构造USRP的端到端链路传输系统，以及写核心器件选型、传感/执行/通信模块、电路原理图关键部分、或实物结构。不要堆料表，重点解释“为什么这样设计”",
        "这里讲如何结合软件和硬件构建平台，实现多样化卫星业务传输场景的构建模拟，以及实际的软硬件协同过程。",
        "这是最拉分的部分。要有测试环境、测试方法、对比对象、指标表格、误差分析。至少准备：\n性能指标表、稳定性测试、极限/异常场景测试、与现有方案对比、失败样例分析。",
        "说明能用在哪里、成本如何、量产/部署难点、后续优化方向",
    ]
    for p in doc.paragraphs:
        if p.text.strip() in markers:
            set_para_text(p, "")


def main():
    shutil.copy2(SRC, OUT)
    doc = Document(OUT)

    remove_instruction_paragraphs(doc)

    # Cover / abstract placeholders
    set_para_text(find_paragraph(doc, "Sparse Code Multiple Access"),
                  "USRP/CUDA LDPC-OFDM Prototype",
                  size_pt=9, bold=False)
    set_para_text(find_paragraph(doc, "Abstract."),
                  "This paper presents a software-hardware co-designed satellite communication verification platform that integrates a multi-standard FEC framework, GPU-accelerated BP-OSD decoding and a UHD C++ USRP transceiver. The prototype uses two NI/Ettus X310 radios connected through PXIe, an OFDM physical layer with QPSK modulation, LDPC coding and a three-stage receiving pipeline. Runtime configuration supports heterogeneous FEC descriptions, including CSR sparse LDPC matrices, 5G NR quasi-cyclic LDPC parameters and Polar reliability sequences. Hardware-in-the-loop video transmission tests show that the stable 4 Msps profile can recover a 985 kB MP4 stream with about 1.25 Mbps measured goodput and sub-1% frame error rate under the current cabled setup, while the UI receives constellation, spectrum and link-metric telemetry through independent rate-limited UDP streams.",
                  size_pt=10.5)
    set_para_text(find_paragraph(doc, "KEY WORDS: key words1, key words2, key words3, key words4"),
                  "KEY WORDS: satellite Internet; USRP; LDPC; GPU BP-OSD; OFDM video transmission; link adaptation",
                  size_pt=10.5)

    # Replace short abstract metric paragraph.
    p_abs = doc.paragraphs[17]
    if "实验结果表明" in p_abs.text:
        p_abs.text = p_abs.text.replace(
            "实验结果表明，所设计系统能够支撑多体制、多业务、多架构条件下的高可靠译码与自适应传输需求，为未来卫星互联网中的通用基带处理、星上智能通信和动态链路优化提供了一种可行技术路径。",
            "实验结果表明，当前原型已完成 UHD 4.9、X310 FPGA 39.3 和 PXIe RIO0/RIO1 条件下的闭环验证；在 4 Msps、512 有效子载波、QPSK、LDPC n=128/k=64 与 GPU BP-OSD 译码配置下，可稳定完成 985 kB MP4 文件恢复，实测 goodput 约 1.25 Mbps，典型 FER 低于 1%，并可向上层 UI 实时输出星座、频谱、SNR、FER 和吞吐率指标。"
        )

    # 3.x system architecture additions
    anchor = find_paragraph(doc, "系统总体架构")
    p = paragraph_after(anchor, "针对当前可运行原型，系统采用“控制端—发送端—接收端”的分层结构。控制端负责统一生成频点、采样率、调制方式、FEC 参数、USRP 设备参数和运行时长；发送端负责媒体切片、帧头封装、LDPC 编码、星座映射与 OFDM 调制；接收端负责 UHD 采样、帧同步、频偏估计、信道估计、星座均衡、GPU/CPU 译码、媒体重组和 UI 遥测输出。")
    p = paragraph_after(p, "底层链路已从 MATLAB/Simulink 原型迁移为 UHD C++/CUDA 实现。USRP 启动、接收采样、同步解调和 GPU 批处理之间通过独立线程与有界队列解耦，避免译码与 UI 显示阻塞 UHD 接收热路径。频谱计算与 UI UDP 发送进一步拆分到限速线程，使星座图、频谱图和 BER/FER 曲线显示不影响主接收链路。")
    table_anchor = paragraph_after(p, "表 1 当前 USRP 闭环验证平台关键参数", style=None)
    table_anchor.alignment = WD_ALIGN_PARAGRAPH.CENTER
    move_table_after(doc, table_anchor, [
        ["项目", "当前配置", "说明"],
        ["射频设备", "NI/Ettus USRP X310 x2，UBX-160", "两台设备通过 PXIe RIO0/RIO1 连接，UHD 4.9.0，FPGA 39.3"],
        ["默认拓扑", "RIO1 RF0 TX/RX -> RIO0 RX", "默认采用 PXIe 设备参数 type=x300,resource=RIO1/RIO0"],
        ["物理层", "OFDM，Nfft=1024，CP=72，48 symbols/frame", "当前调制为 QPSK，支持 activeSC 与 pilot 周期配置"],
        ["稳定演示档", "4 Msps，activeSC=512，pilotPeriod=4", "当前视频传输与 UI 展示的推荐参数"],
        ["FEC/译码", "LDPC n=128/k=64，GPU BP-OSD；CPU 译码可切换", "短码用于 GPU BP-OSD 闭环演示，接口预留长码 LDPC/Polar"],
        ["上层遥测", "UDP 65432/65433/65434", "分别发送星座点、频谱和链路指标"],
    ])

    # 4.1 FEC section
    anchor = find_paragraph(doc, "多体制兼容的通用信道译码算法")
    p = paragraph_after(anchor, "当前工程已将 FEC 参数从硬编码方式改为动态配置方式。为避免不同标准之间的数据结构被错误统一，系统将编码体制划分为三类：DVB-S2/CCSDS 类长稀疏 LDPC 采用 CSR 稀疏矩阵；5G NR LDPC 采用 Base Graph 与提升因子 Zc 的准循环参数；Polar 码采用 N、K 与可靠性序列 Q-sequence。MatrixLoader 在启动时读取对应参数并生成统一 FecConfig，TX/RX 仅依赖抽象接口，不再直接绑定固定矩阵。")
    p = paragraph_after(p, "RX 初始化阶段根据 FecConfig 执行严格路由：当 block_length N 大于 1000 时强制进入 GPU BP 译码路径；短 Polar 码可进入 SCL 译码；短 LDPC 码可进入 CPU min-sum 或显式 GPU BP-OSD 路径。该策略保证 DVB-S2 normal frame 等长码字不会误走低吞吐 CPU 矩阵消元路径，也为后续 5G NR QC-LDPC 专用 kernel 接入保留边界。")
    table_anchor = paragraph_after(p, "表 2 编码体制与运行时数据结构", style=None)
    table_anchor.alignment = WD_ALIGN_PARAGRAPH.CENTER
    move_table_after(doc, table_anchor, [
        ["类型", "代表体制", "参数结构", "当前状态"],
        ["Type A", "DVB-S2/CCSDS LDPC", "CSR(row_ptr, col_ind, val)，支持 cudaMemcpy", "n128/k64、n512/k256 完成闭环；DVB-S2 n64800/k32400 完成加载和路由验证"],
        ["Type B", "5G NR LDPC", "BG1/BG2 + Zc，不展开完整矩阵", "完成参数结构和 GPU BP 强制路由验证"],
        ["Type C", "Polar", "N、K、reliability sequence", "完成可靠性序列加载和短码 SCL 路由验证"],
    ])

    # GPU/FPGA implementation section
    anchor = find_paragraph(doc, "基于GPU架构的通用译码并行实现")
    p = paragraph_after(anchor, "GPU 译码路径当前以 BP-OSD 短码译码器为核心验证对象，配置为 LDPC n=128、k=64，支持最小 batch、最大 batch 和延迟上限参数。接收端将 OFDM 解调后的软信息按码块聚合后送入 GPU 批处理队列，GPU kernel 完成 BP 迭代与 OSD 辅助判决，译码结果再回传媒体重组模块。")
    p = paragraph_after(p, "为降低实时链路抖动，工程将 UHD 接收、OFDM 同步/解调、GPU batch 调度和 UI 发送拆成多个异步阶段。UHD 线程只负责连续收样和入队；基带线程完成下采样、同步峰搜索、CFO 估计、信道估计和均衡；译码线程负责 FEC 解码与 CRC 校验；UI 线程按照固定周期发送抽样星座、PSD 和指标。该设计使接收端在 UI 打开时仍能保持稳定吞吐。")

    # 5.2 USRP section
    anchor = find_paragraph(doc, "基于USRP的星地链路收发系统实现")
    p = paragraph_after(anchor, "USRP 收发机采用独立启动模块封装，发送端和接收端均通过 UHD device args 指定 PXIe 资源。默认配置为 RIO1 作为发送 USRP、RIO0 作为接收 USRP，主时钟 200 MHz，推荐中心频率 5 GHz，采样率 4 Msps，TX/RX 口用于发射，RX2 或指定 RF 通道用于接收。脚本层提供 Profile、频点、增益、天线、activeSC、pilotPeriod、同步阈值、队列长度和 GPU 译码器等参数入口，便于上层控制端直接下发。")
    p = paragraph_after(p, "当前物理层帧结构保持固定：每帧包含重复 QPSK 前导用于 Schmidl 类同步和 CFO 估计，随后为 48 个 OFDM 符号；在频域内使用 1024 点 FFT、72 点循环前缀和可配置有效子载波。块导频按 pilotPeriod 周期嵌入 OFDM 符号，用于信道估计和数据子载波均衡。媒体业务通过 48 字节帧头携带 stream id、chunk id、文件大小和 CRC，接收端只接受通过 FEC 与 CRC 校验的 chunk。")
    table_anchor = paragraph_after(p, "表 3 当前推荐视频演示参数", style=None)
    table_anchor.alignment = WD_ALIGN_PARAGRAPH.CENTER
    move_table_after(doc, table_anchor, [
        ["参数", "推荐值", "原因"],
        ["连接方式", "PXIe，TX=RIO1，RX=RIO0", "避免网口 MTU 和驱动路径差异，UHD 路径通过 UHD_ROOT 指定"],
        ["主频", "5 GHz", "实测比 2.45 GHz 更适合当前实验室连线环境"],
        ["采样率", "4 Msps", "在吞吐与 FER 之间取得当前最优平衡"],
        ["activeSC/pilot", "512 / 4", "保证帧净荷与同步/信道估计稳定性"],
        ["增益", "TxGain 约 12-20 dB，RxGain 约 25-30 dB", "根据端口和衰减器微调，避免近端过驱和低 SNR 突降"],
        ["同步阈值", "0.50 左右", "降低误检，同时保持帧捕获能力"],
        ["译码器", "cuda-bp-osd 或 cpu", "GPU 用于 n128/k64 演示，CPU 用于对照测试"],
    ])

    # Test section
    anchor = find_paragraph(doc, "6.1 测试环境与评价指标")
    p = paragraph_after(anchor, "当前硬件在 Windows 平台运行，UHD 版本为 4.9.0.0-release，USRP X310 固件版本 6.1，FPGA 版本 39.3，设备通过 NI-RIO PXIe 资源 RIO0/RIO1 访问。测试视频包括 test_video1.mp4，大小 985007 bytes。评价指标包括同步峰值、CFO、SNR、FER、goodput、文件 chunk 完成度、GPU batch 延迟和 UI 显示连续性。")
    p = paragraph_after(p, "测试流程为：先启动 RX 并预留 lead time，再启动 TX；TX 循环发送视频 chunk，RX 按 chunk id 去重并重组文件；同时 UI 端通过 UDP 接收星座点、频谱和指标曲线。除链路成功率外，测试还记录不同采样率、不同 activeSC、不同收发增益和不同 RF 端口下的 FER 变化，用于确定演示 profile。")
    table_anchor = paragraph_after(p, "表 4 当前 USRP 视频传输实测结果摘要", style=None)
    table_anchor.alignment = WD_ALIGN_PARAGRAPH.CENTER
    move_table_after(doc, table_anchor, [
        ["测试项", "参数", "结果"],
        ["离线 FEC 架构验证", "CSR LDPC、NR QC-LDPC、Polar Q-sequence", "全部通过；N>1000 自动路由 GPU BP"],
        ["离线 OFDM/LDPC 闭环", "n512/k256，activeSC=512", "2 帧全通过，FER=0，BER=0"],
        ["2 Msps 视频链路", "activeSC=512，UI 开启", "文件完整恢复，goodput 约 0.61-0.62 Mbps"],
        ["4 Msps 视频链路", "activeSC=512，UI 开启", "文件完整恢复，典型 FER 约 0.2%-0.8%，goodput 约 1.25 Mbps"],
        ["8 Msps 扫参", "activeSC=256/512", "存在同步和低 SNR 突降，短测未优于 4 Msps"],
        ["RIO0 RF1 TX/RX 接收口", "5 GHz，RxGain 25", "可捕获同步峰，但 SNR 约 12 dB，FER 约 20%，不作为默认演示口"],
    ])

    # Subsections
    for heading, text in [
        ("高可靠通用信道译码性能测试", "离线验证已经覆盖三类 FEC 参数结构。CCSDS n128/k64 和 n512/k256 可完成编码、OFDM 调制、解调与译码闭环；DVB-S2 n64800/k32400 已完成 alist 生成、CSR 加载和 GPU BP 强制路由验证。受限于当前旧 CPU 编码路径对长矩阵系统化消元开销较大，DVB-S2 normal frame 的完整端到端链路下一步将接入专用 IRA/QC 编码器和通用 CUDA BP backend。"),
        ("GPU软硬件协同译码性能测试", "GPU BP-OSD 译码器已接入 USRP 视频链路，可在 RX 端通过 --decoder cuda-bp-osd 启用。接收端按 minBatch、maxBatch 和 latencyUs 调度 GPU batch，并将译码结果交给媒体 CRC 与 chunk 重组模块。CPU 译码路径保留为可靠性对照，用于区分 RF 问题、同步问题和 GPU 译码问题。"),
        ("预测增强链路自适应传输性能测试", "当前工程已实现基于 RX 反馈的自适应重复传输机制。RX 按滑动窗口统计平均 SNR 与 FER，通过 UDP feedback 将推荐 repeat 发送给 TX；当 FER 超过阈值或 SNR 下降时，TX 自动提高每帧重复次数，优先保障文件恢复可靠性；当链路恢复时再降低重复次数以提升吞吐。该机制已在 8 Msps 不稳定链路中观测到 repeat 由 1 自动切换到 3。"),
        ("虚实融合端到端系统验证", "端到端验证已完成控制端、发送端、接收端和 UI 的数据边界定义。控制端负责下发 profile、频点、采样率、增益、FEC 与媒体文件；发送端向 UI 输出理想星座点；接收端向 UI 输出均衡星座点、频谱 PSD、SNR、FER、goodput 和接收文件路径。UI 数据发送从 decode_frame 热路径移出后，显示刷新不会明显拖慢链路解调。"),
    ]:
        try:
            anchor = find_paragraph(doc, heading)
            paragraph_after(anchor, text)
        except ValueError:
            pass

    # Final summary refinement
    try:
        anchor = find_paragraph(doc, "作品总结与未来工作")
        paragraph_after(anchor, "围绕 USRP 与编译码闭环，本文当前已形成可运行的软件无线电样机：一方面通过 MatrixLoader 和 FecConfig 将 LDPC/Polar 等不同编码体制抽象为运行时可切换配置；另一方面通过 UHD C++/CUDA 构建两台 X310 的 OFDM 视频传输链路，并完成 GPU BP-OSD、CPU 译码、UI 遥测和自适应重复传输的联调。后续重点将从短码演示扩展到 DVB-S2 长码专用编码器和通用 CUDA BP backend，并继续提升 8 Msps 以上链路的同步稳定性与有效吞吐。")
    except ValueError:
        pass

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
